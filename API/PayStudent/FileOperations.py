#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
# @Time   : 2024/06/28 18:11
# @Author : 新贵大人
描述:
"""
import os
import openpyxl
from datetime import datetime

class FileHandler:
    def __init__(self, environment, subject):
        self.environment = environment
        self.subject = subject

    def get_path(self, file_name):
        file_local_path_name = os.path.join(os.path.dirname(os.path.realpath(__file__)), f"{file_name}.xlsx")
        return file_local_path_name

    # 获取系统没有重复使用的手机号
    def get_student_phone(self) -> list:
        # 检查结果是否已经缓存
        if self.cached_phone_list is not None:
            return self.cached_phone_list
        phone_list = []
        # 判断是否需要自定义手机号
        if self.custom_number == "0":
            try:
                for i in range(self.num):
                    now = time.time()
                    new_phone = '1' + str(now).replace('.', '')[-10:]
                    url = f'https://{self.environment}-gw.vipthink.cn/api/member/v3/back/ol-user/getUserInfoByMobile'
                    payload = {"mobile": f"{new_phone}"}
                    resp = self.handle_api_response(
                        requests.post(url, json=payload, headers={"authorization": self.token}))
                    if resp and resp.get('code') == 0 and not resp.get('data'):
                        if phone_list != []:
                            # 判断列表中有没有相同的手机号码
                            if new_phone in [i for i in phone_list]:
                                if int(new_phone[-1:]) != 9:
                                    # \d表示匹配一个数字字符，$表示匹配字符串的结尾位置
                                    phone_list.append(re.sub(r'\d$', str(int(new_phone[-1:]) + 1), new_phone))
                                else:
                                    phone_list.append(re.sub(r'\d$', str(int(new_phone[-1:]) - 1), new_phone))
                            else:
                                phone_list.append(new_phone)
                        else:
                            phone_list.append(new_phone)
            except TypeError as error:
                print('API request exception, please check\n', 'error:', error)
            self.cached_phone_list = phone_list
            return phone_list
        else:
            return self.custom_number

    def operation_table(self, data_list):
        formatted_datetime1 = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        formatted_datetime2 = datetime.now().strftime("%Y%m%d%H%M%S")
        environment_settings = {
            "uat": {
                "kc_package_skuId": '31825515',
                "kc_order_amount": '0.02',
                "mmears_package_skuId": '31830439',
                "mmears_order_amount": '0.43',
                "yz_package_skuId": '10025804',
                "yz_order_amount": '2.3'
            },
            "preprod": {
                "kc_package_skuId": '20528464',
                "kc_order_amount": '0.01',
                "yz_package_skuId": '10016610',
                "yz_order_amount": '0.01'
            }
        }
        settings = environment_settings[self.environment]
        file_name = 'biaoda' if self.subject in ["口才", "魔力"] else 'yizhi'
        file_path = self.get_path(file_name)
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook["Sheet1"] if self.subject in ["口才", "魔力"] else workbook["导入主表"]

        if sheet.max_row > 2:
            for row in range(3, sheet.max_row + 1):
                sheet.delete_rows(row)

        for index, value in enumerate(data_list):
            sheet.cell(row=3 + index, column=2, value='0')
            sheet.cell(row=3 + index, column=4, value="86")
            sheet.cell(row=3 + index, column=10, value=formatted_datetime1)
            sheet.cell(row=3 + index, column=11, value="free")
            sheet.cell(row=3 + index, column=12, value="0")
            sheet.cell(row=3 + index, column=13, value="1999")
            sheet.cell(row=3 + index, column=15, value="0")
            sheet.cell(row=3 + index, column=1, value=f'XG{formatted_datetime2}{value}')
            sheet.cell(row=3 + index, column=5, value=value)
            if self.subject == "口才":
                sheet.cell(row=3 + index, column=8, value=settings["kc_package_skuId"])
                sheet.cell(row=3 + index, column=9, value=settings["kc_order_amount"])
            elif self.subject == "魔力":
                sheet.cell(row=3 + index, column=8, value=settings["mmears_package_skuId"])
                sheet.cell(row=3 + index, column=9, value=settings["mmears_order_amount"])
            elif self.subject == "益智":
                sheet.cell(row=3 + index, column=6, value="468078")
                sheet.cell(row=3 + index, column=7, value=settings["yz_package_skuId"])
                sheet.cell(row=3 + index, column=8, value=settings["yz_order_amount"])
                sheet.cell(row=3 + index, column=9, value=formatted_datetime1)
                sheet.cell(row=3 + index, column=10, value="第三方售卖")
                sheet.cell(row=3 + index, column=11, value="0")
                sheet.cell(row=3 + index, column=12, value="1999")
                sheet.cell(row=3 + index, column=13, value="")
        workbook.save(file_path)
        workbook.close()
