#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
# @Time   : 2024/04/16 15:56
# @Author : 新贵大人
描述:
"""
import json
import os

import pymysql
import datetime as dt
import requests
from openpyxl import load_workbook

mysql_conn_preprod_manager = pymysql.connect(host='db.preprod.61draw.com', port=3306, user='root', password='dbtest',
                                             db='i61-hll-manager')
mysql_conn_preprod_i61 = pymysql.connect(host='db.preprod.61draw.com', port=3306, user='root', password='dbtest',
                                         db='i61')
mysql_conn_test_manager = pymysql.connect(host='testdb.61info.com', port=3306, user='root', password='dbtest',
                                          db='i61-hll-manager')
mysql_conn_test_i61 = pymysql.connect(host='testdb.61info.com', port=3306, user='root', password='dbtest', db='i61')

def get_benefit_card_id(user_id):
    cursor = mysql_conn_test_manager.cursor()
    sql = f"select * from user_benefit_card_record where user_id = {user_id} and state = 0 limit 1;"
    cursor.execute(sql)
    results = cursor.fetchall()[0]
    return results[0]

# 更新减少权益卡文件参数
def reduce_xg_file(apk,user_id) -> dict:
    wb = load_workbook(apk)
    ws = wb.active

    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    try:
        sys_code_col = header.index("系统生成编码ID") + 1   # openpyxl 列号从 1 开始
    except ValueError:
        raise ValueError("Excel 中未找到“系统生成编码ID”列，请检查表头")

    for row in ws.iter_rows(min_row=2, values_only=False):
        new_code = get_benefit_card_id(user_id)
        row[sys_code_col - 1].value = new_code

    wb.save(apk)
    with open(apk, 'rb') as f:
        file_stream = f.read()

    file_name = os.path.basename(apk)
    return {"type": 2, "file": {"name": file_name,"content": file_stream}}

# 获取当前批次，如果存在则删除
def delete_batch():
    cursor = mysql_conn_preprod_manager.cursor()
    sql = "SELECT * FROM allocation_batch WHERE type = 3 AND NOW() BETWEEN begin_time AND end_time;"
    try:
        cursor.execute(sql)
        results = cursor.fetchall()
        if results:
            for row in results:
                batch_id = row[0]
                delete_sql = f"DELETE FROM allocation_batch WHERE id = {batch_id};"
                cursor.execute(delete_sql)
                mysql_conn_preprod_manager.commit()
                # print("批次:{},删除成功".format(batch_id))
                return 1
        else:
            print("未找到匹配的批次")
            return 1
    except Exception as e:
        return f"查询出错：{str(e)}"
    finally:
        cursor.close()
        mysql_conn_preprod_manager.close()


# 删除对应强化课分班的批次
def delete_enhance_batch():
    cursor = mysql_conn_preprod_manager.cursor()
    sql = "SELECT * FROM allocation_batch WHERE type = 4 AND end_time >= NOW();"
    try:
        cursor.execute(sql)
        results = cursor.fetchall()
        if results:
            for row in results:
                batch_id = row[0]

                delete_sql = f"DELETE FROM allocation_batch WHERE id = {batch_id};"
                cursor.execute(delete_sql)
                delete_related_sql = f"DELETE FROM allocation_batch_detail WHERE batch_id = {batch_id};"
                cursor.execute(delete_related_sql)

                mysql_conn_preprod_manager.commit()
                # print("批次:{},删除成功".format(batch_id))
                return 4
        else:
            print("未找到匹配的批次")
            return 4
    except Exception as e:
        return f"查询出错 {str(e)}"
    finally:
        cursor.close()
        mysql_conn_preprod_manager.close()



# 获取level的学员id，可用于提交升阶组班意向的
def get_level_one_student(sort):
    cursor1 = mysql_conn_preprod_manager.cursor()
    cursor2 = mysql_conn_preprod_i61.cursor()

    try:
        sql1 = """
            SELECT DISTINCT ugr.user_id
            FROM `i61-hll-manager`.user_group_relation ugr
            JOIN `i61-hll-manager`.group_info gi ON ugr.group_id = gi.id
            JOIN `i61-hll-manager`.allocation_application aa ON ugr.user_id = aa.user_id
            JOIN `i61-hll-manager`.group_table_relation gtr ON ugr.group_id = gtr.group_id AND gtr.state = 0
            JOIN `i61-hll-manager`.course_table ct ON gtr.table_id = ct.id
            JOIN `i61-hll-manager`.course_begin_stage_info cbsi ON ct.begin_course_stage_id = cbsi.id
            LEFT JOIN `i61-hll-manager`.allocation_application_advance aaa ON aaa.user_id = ugr.user_id
            WHERE cbsi.years = 1
            AND ugr.group_id > 0
            AND gtr.state = 0
            AND (aaa.state = 1 OR aaa.id IS NULL)
            AND aa.state = 7
            ORDER BY RAND() DESC
            LIMIT 30;
        """

        cursor1.execute(sql1)
        student_ids = [row[0] for row in cursor1.fetchall()]
        student_ids_str = ','.join(map(str, student_ids))

        if not student_ids:
            return []

        sql2 = f"""
            SELECT UserId
            FROM i61.usercoursepackageinfo cpi
            WHERE cpi.UserId IN ({student_ids_str}) AND cpi.State < 4
            GROUP BY UserId
            HAVING SUM(cpi.CourseNumber) + SUM(cpi.GiftCourseNumber) - SUM(cpi.OpenCourseNumber) - SUM(cpi.OpenGiftCourseNumber) > 0
            LIMIT 3;
        """

        cursor2.execute(sql2)
        selected_students = [row[0] for row in cursor2.fetchall()]

        if not selected_students:
            return []

        sql3 = f"""
            SELECT * FROM `i61-hll-manager`.allocation_application_old
            WHERE user_id = {selected_students[sort - 1]};
        """
        cursor1.execute(sql3)
        results3 = cursor1.fetchall()

        if results3:
            sql4 = f"""
                DELETE FROM `i61-hll-manager`.allocation_application_old
                WHERE user_id = {selected_students[sort - 1]};
            """
            cursor1.execute(sql4)
            mysql_conn_preprod_manager.commit()

        cursor1.close()
        cursor2.close()
        mysql_conn_preprod_manager.close()
        mysql_conn_preprod_i61.close()

        if sort in [1, 2, 3]:
            return selected_students[sort - 1] if len(selected_students) >= sort else None
        return selected_students

    except Exception as e:
        print(f"发生错误: {e}")
        cursor1.close()
        cursor2.close()
        mysql_conn_preprod_manager.close()
        mysql_conn_preprod_i61.close()


def get_week_day(weekday):
    # 获取当前日期
    current_date = dt.datetime.now().date()
    # 获取当前日期是星期几（0=周一，1=周二，...，6=周日）
    current_weekday = current_date.weekday()

    # 计算距离目标星期的天数
    days_until_target = weekday - current_weekday - 1
    if days_until_target <= 0:
        days_until_target += 7

    # 计算最近的目标日期
    nearest_target_date = current_date + dt.timedelta(days=days_until_target)
    # 格式化日期
    nearest_target_date_formatted = nearest_target_date.strftime("%Y-%m-%d")
    return nearest_target_date_formatted



def future_date(num_day):
    today = dt.datetime.now().date()
    dates_list = []
    for i in range(num_day):
        future_date = today + dt.timedelta(days=i)
        dates_list.append(future_date)
    return dates_list


# 获取15.20-16.30分的批次详情id，用于开放时段
def get_batch_detailId(token,batch_id, host = "https://gw-mg-preprod.61info.cn", time_id = 96):
    url = f"{host}/manager-api/o/new/applyAllocate/getBatchDetail.json"
    headers = {"authorization": token, 'content-type': 'application/x-www-form-urlencoded;charset=UTF-8'}
    payload = f'batchId={batch_id}&needRecommendGroupCount=0'
    resp = requests.post(url, data=payload, headers=headers, timeout=5).json()
    # 提取timeId为96的数据
    time_data = next((item for item in resp['data']['dataList'] if item['timeId'] == time_id), None)
    # 如果找到了timeId的数据，提取周一的detailId
    if time_data:
        detail_id = next((vo['detailId'] for vo in time_data['voList'] if vo['dayForWeek'] == 1), None)
        return detail_id


# 修改强化课批次的分班结束时间为当天，不然强化课分班不成功
def enhance_batch_endtime(batch_id):
    try:
        today = dt.datetime.now().date()
        end_of_today = dt.datetime.combine(today, dt.time(23, 59, 59))
        cursor = mysql_conn_preprod_manager.cursor()

        update_sql1 = """
            UPDATE `i61-hll-manager`.allocation_batch
            SET end_time = %s
            WHERE id = %s;
        """
        cursor.execute(update_sql1, (end_of_today, batch_id))

        update_sql2 = """
            UPDATE `i61-hll-manager`.allocation_batch_consequence_detail
            SET allocation_date = %s
            WHERE batch_id = %s;
        """
        cursor.execute(update_sql2, (end_of_today, batch_id))

        mysql_conn_preprod_manager.commit()

    except Exception as e:
        mysql_conn_preprod_manager.rollback()
        return f"更新失败：{str(e)}"

    finally:
        cursor.close()
        mysql_conn_preprod_manager.close()

    return 1

# 更新强化课学员的已选分班状态
def update_allocation_application_enhance_state(student_id):
    try:
        cursor = mysql_conn_preprod_manager.cursor()
        # 使用参数化查询来防止SQL注入
        update_sql = """
            UPDATE `i61-hll-manager`.allocation_application_enhance
            SET state = 2
            WHERE user_id = %s;
        """
        cursor.execute(update_sql, student_id)
        mysql_conn_preprod_manager.commit()
    except Exception as e:
        return f"更新失败：{str(e)}"
    finally:
        cursor.close()
        mysql_conn_preprod_manager.close()
    return 1

# 获取当前时间
def today_date():
    today = dt.datetime.now().date()
    return str(today)

# 获取n天后的时间
def future_num_date(num_day):
    future_date = dt.datetime.now().date() + dt.timedelta(days=num_day)
    return str(future_date)

if __name__ == '__main__':
    # print(get_batch_detailId("eyJhbGciOiJIUzI1NiJ9.eyJkYXRhIjoicGgzMk4wK0diSThPZEhNMnZpbUJ5WDlKM2N2RXIxbFl0YVhZdmVxd1dENXV4bzJ0bCt2d0NkalJDTlBLR2lueG1yT2NiZlJuVys1eCttRnF3eTdBVFlIMEdWb0lEdk4vZGxLL3o1NXdvS0hJVFZJNHEwTUF4SEpLWkNISW1lUTNiY1dtUk9VSjlrMzZNOGpaVWV6SlFrMHhRT0d6OE5aNmFwTnpuYys4LzZBVGVoc2lXUVMybmlKVWVYWThNMDhMSDU5aXA2UWNjMVM4SUxNY2VDK2YzUnRGSjNsKzZ5eHNIVm9vY2pkNWNpeGloUkZBK0RVdEhnM0dWZDJrM2M3bXFBZ1h5Wld5TG4rUHRoZ0dYT1ZsaUJlRXp4Y1FQRDhFV1pWT2lZcTFHOW89IiwiZXhwIjoxNzY4MDMzNDcxfQ.UMUxzEpeh6dQIDj9C_lh-uPxQ0_uV6N0ivmILRttDqI",1216))
    # print(get_benefit_card_id(22562397))
    print(reduce_xg_file(r"C:\Users\t\Desktop\aaa.xlsx","22562397"))